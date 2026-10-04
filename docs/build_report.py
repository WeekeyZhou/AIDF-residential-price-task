"""Build 'Summary Report.docx' from summary_report.md (needs pandoc and python-docx)."""
import subprocess
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = "Summary Report.docx"
subprocess.run(["pandoc", "summary_report.md", "-o", OUT], check=True)
d = Document(OUT)
for s in d.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
for st in d.styles:
    if st.type != 1:
        continue
    if st.name.startswith("Heading"):
        st.font.size = Pt(12.5)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(8), Pt(3)
    elif st.name == "Title":
        st.font.size = Pt(16)
    elif st.name == "Source Code":
        st.font.size = Pt(8.5)
    else:
        st.font.size = Pt(10)
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(0), Pt(4)
for t in d.tables:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "999999")):
            el.set(qn(f"w:{k}"), v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
                    r.font.bold = True if i == 0 else r.font.bold
for p in d.paragraphs:
    if p.style.name == "Source Code":
        for r in p.runs:
            r.font.size = Pt(8.5)
d.save(OUT)
